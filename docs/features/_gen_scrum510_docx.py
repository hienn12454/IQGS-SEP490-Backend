"""SCRUM-510: tách question regen khỏi Generate Unlimited + làm rõ admin HR limits."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt

OUT = Path(r"e:\Github\SEP490\code\docs\features")
OUT.mkdir(parents=True, exist_ok=True)
PATH = OUT / "scrum510-hr-question-regen-admin-limits.docx"


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)
    return p


def main():
    doc = Document()
    add_para(doc, "1. THÔNG TIN CHUNG", True)
    add_para(
        doc,
        "Tên chức năng : Tách hạn mức regen câu HR khỏi Unlimited tạo bộ + làm rõ UI admin / "
        "Decouple HR question regen from generateUnlimited + clarify admin plan limits",
    )
    add_para(doc, "Jira Ticket : SCRUM-510")
    add_para(doc, f"Ngày tạo : {date.today().strftime('%d/%m/%Y')}")
    add_para(doc, "Người thực hiện: nhóm IQGS")
    add_para(doc, "Trạng thái : Done")

    add_para(doc, "2. VỊ TRÍ TRONG HỆ THỐNG", True)
    add_para(doc, "Layer : Domain / Application / FE Admin + HR Settings")
    add_para(doc, "Module : Shared — Subscription / Studio HR")
    add_para(
        doc,
        "Files: SubscriptionGateService.cs, HrGenerateWindowTests.cs, SubscriptionPlanLimits.cs, "
        "hr-usage-panel.tsx, error.interceptor.ts, admin-plans-page.tsx, vi.ts, en.ts",
    )

    add_para(doc, "3. CÔNG NGHỆ & KỸ THUẬT", True)
    add_para(
        doc,
        "- Feature gate đọc LimitsSnapshotJson; UsageCounter HrQuestionRegen scope = InterviewPlanId (N). "
        "Không migration.",
    )
    add_para(
        doc,
        "- Admin UI nhóm 3 hạn mức Studio (tạo bộ/JD, regen từng câu, refine outline) để tránh nhầm "
        "Generate Unlimited với regen câu.",
    )

    add_para(doc, "4. FLOW MÔ TẢ", True)
    add_para(
        doc,
        "1) Admin đặt questionRegenPerPlan = N (N>0) trên Free hoặc Premium. "
        "2) HR regen từng câu trong Studio → CheckQuestionRegenAsync so used >= N. "
        "3) GenerateUnlimited chỉ skip CheckGenerateSetAsync; không skip regen câu. "
        "4) Cửa sổ tạo bộ: used < max → cho qua dù LastSuccessfulGenerateAt còn (sau khi admin tăng N).",
    )

    add_para(doc, "5. LÝ DO THIẾT KẾ", True)
    add_para(
        doc,
        "a) UI admin đã cho sửa regen khi Unlimited bật → BE phải enforce field đó độc lập. "
        "b) Đã cân nhắc gộp Unlimited = mọi AI unlimited — quá rộng, trái copy admin. "
        "c) Trade-off: Premium mặc định QuestionRegenPerPlan=0 vẫn unlimited cho đến khi Admin đặt số > 0.",
    )

    add_para(doc, "6. GHI CHÚ & GIỚI HẠN", True)
    add_para(doc, "- Studio không có regen cả bộ; counter là tổng lần regen câu trên 1 plan.")
    add_para(doc, "- Sync limits Active không reset UsageCounter / Last.")
    add_para(doc, "- Toast VI ưu tiên detail BE (số max động); EN dùng fallback interceptor.")

    doc.save(PATH)
    print(f"Wrote {PATH}")


if __name__ == "__main__":
    main()
